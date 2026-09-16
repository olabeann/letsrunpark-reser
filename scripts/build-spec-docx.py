"""Sync the editable specification Markdown to its existing Word deliverable."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'outputs/렛츠런파크_기능명세서_v1.md'
TARGET = SOURCE.with_suffix('.docx')

def plain(text):
    return re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', text).replace('**', '').replace('`', '')

def build():
    doc = Document(TARGET) if TARGET.exists() else Document()
    for child in list(doc._element.body):
        if child.tag != qn('w:sectPr'):
            doc._element.body.remove(child)
    for name, size in [('Normal',11),('Title',24),('Subtitle',11),('Heading 1',16),('Heading 2',12),('Heading 3',11),('List Bullet',11)]:
        style = doc.styles[name]
        style.font.name = '맑은 고딕'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0,0,0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '맑은 고딕')
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.18
        if name.startswith('Heading'):
            style.paragraph_format.space_before = Pt(14 if name == 'Heading 1' else 10)
            style.paragraph_format.keep_with_next = True
        borders = style.element.xpath('./w:pPr/w:pBdr')
        for el in borders: el.getparent().remove(el)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                p.clear()
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    pos=0
    while pos < len(lines):
        line=lines[pos].strip()
        if not line:
            pos+=1; continue
        if line.startswith('|'):
            rows=[]
            while pos<len(lines) and lines[pos].strip().startswith('|'):
                cells=[plain(x.strip()) for x in lines[pos].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x) for x in cells): rows.append(cells)
                pos+=1
            table=doc.add_table(rows=1,cols=len(rows[0]))
            table.autofit=False
            width=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
            two_column_widths = {'상황': [.38,.62], '역할': [.23,.77], 'ID': [.13,.87]}
            fractions=two_column_widths.get(rows[0][0],[.3,.7]) if len(rows[0])==2 else ([.13,.43,.44] if len(rows[0])==3 else [1/len(rows[0])]*len(rows[0]))
            for col,fraction in zip(table.columns,fractions): col.width=int(width*fraction)
            pr=table._tbl.tblPr
            for el in pr.findall(qn('w:tblBorders')): pr.remove(el)
            borders=OxmlElement('w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                item=OxmlElement('w:'+side); item.set(qn('w:val'),'single'); item.set(qn('w:sz'),'4'); item.set(qn('w:color'),'D9D9D9'); borders.append(item)
            pr.append(borders)
            for i,row_values in enumerate(rows):
                cells=table.rows[0].cells if i==0 else table.add_row().cells
                row=table.rows[i]
                trpr=row._tr.get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if i==0: trpr.append(OxmlElement('w:tblHeader'))
                for j,(cell,value) in enumerate(zip(cells,row_values)):
                    cell.width=int(width*fractions[j])
                    cell.text=value
                    tcpr=cell._tc.get_or_add_tcPr()
                    margins=OxmlElement('w:tcMar')
                    for side in ('top','left','bottom','right'):
                        m=OxmlElement('w:'+side); m.set(qn('w:w'),'100'); m.set(qn('w:type'),'dxa'); margins.append(m)
                    tcpr.append(margins)
                    if i==0:
                        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EDEDED'); tcpr.append(shade)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(2)
                        p.paragraph_format.keep_with_next=i==0
                        for run in p.runs:
                            run.font.size=Pt(10)
                            run.bold=i==0
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
            continue
        if line.startswith('# '):
            p=doc.add_paragraph(plain(line[2:]),'Title')
        elif line.startswith('## '):
            p=doc.add_paragraph(plain(line[3:]),'Heading 1')
        elif line.startswith('### '):
            heading=plain(line[4:])
            heading=re.sub(r'^([A-Z]+)-(\d+)',r'\1 \2',heading)
            p=doc.add_paragraph(heading,'Heading 2')
        elif line.startswith('- '):
            p=doc.add_paragraph(plain(line[2:]),'List Bullet')
        else:
            p=doc.add_paragraph(plain(line))
        p.paragraph_format.widow_control=True
        pos+=1
    doc.core_properties.title='렛츠런파크 통합 예약 결제 시스템 기능명세서'
    doc.core_properties.subject='2026년 9월 16일 화면 기준 정책 개정 v1.1'
    doc.core_properties.comments=''
    doc.save(TARGET)
    check=Document(TARGET)
    all_text='\n'.join(p.text for p in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for r in t.rows for c in r.cells)
    for term in ('v1.1','S-01','S-05','AT-19','다음 달 8일','1명 카드'):
        assert term in all_text, term
    assert len(check.tables)==6, len(check.tables)
    print(f'Built {TARGET.name}: {len(check.paragraphs)} paragraphs, {len(check.tables)} tables')

if __name__ == '__main__': build()
